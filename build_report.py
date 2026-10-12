import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # 1. Page Margins Setup (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Palette Constants
    COLOR_PRIMARY_NAVY = RGBColor(27, 54, 93)      # #1B365D - Deep Corporate Navy
    COLOR_SECONDARY_BLUE = RGBColor(46, 91, 136)   # #2E5B88 - Slate Blue
    COLOR_ACCENT_TEAL = RGBColor(14, 116, 144)     # #0E7490 - Modern Accent
    COLOR_CHARCOAL = RGBColor(34, 34, 34)          # #222222 - Crisp readable body
    COLOR_MUTED = RGBColor(100, 116, 139)          # #64748B - Muted gray
    HEX_HEADER_BG = "1B365D"
    HEX_LIGHT_ROW = "F8FAFC"
    HEX_CODE_BG = "F1F5F9"
    HEX_CONSOLE_BG = "0F172A"                       # Dark slate console
    HEX_BORDER = "CBD5E1"

    # Style Configurations
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_CHARCOAL
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1", sz="4"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_custom_heading(text, level):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        if level == 1:
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            run.font.size = Pt(18)
            run.font.color.rgb = COLOR_PRIMARY_NAVY
        elif level == 2:
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(14)
            run.font.color.rgb = COLOR_SECONDARY_BLUE
        elif level == 3:
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_ACCENT_TEAL
        return h

    def add_callout(text, label="PLACEMENT & INDUSTRY INSIGHT"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        set_cell_background(cell, "EFF6FF") # Light blue tint
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        # Left border highlight
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run(f"[{label}] ")
        r_lbl.bold = True
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(37, 99, 235)

        r_txt = p.add_run(text)
        r_txt.font.size = Pt(10.5)
        r_txt.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph() # Spacing

    def add_code_block(code_text, filename=""):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="0" w:color="0E7490"/><w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2)

        if filename:
            r_fn = p.add_run(f"// File: {filename}\n")
            r_fn.font.name = 'Consolas'
            r_fn.font.size = Pt(9)
            r_fn.bold = True
            r_fn.font.color.rgb = RGBColor(14, 116, 144)

        r_code = p.add_run(code_text)
        r_code.font.name = 'Consolas'
        r_code.font.size = Pt(9)
        r_code.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph() # Spacing

    def add_terminal_output(output_text, title="Verified Execution Output (Terminal)"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        set_cell_background(cell, HEX_CONSOLE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="0" w:color="22C55E"/><w:top w:val="single" w:sz="4" w:space="0" w:color="334155"/><w:right w:val="single" w:sz="4" w:space="0" w:color="334155"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="334155"/></w:tcBorders>')
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2)

        r_hdr = p.add_run(f"$ {title}\n")
        r_hdr.font.name = 'Consolas'
        r_hdr.font.size = Pt(8.5)
        r_hdr.bold = True
        r_hdr.font.color.rgb = RGBColor(34, 197, 94) # Green prompt

        r_out = p.add_run(output_text.strip())
        r_out.font.name = 'Consolas'
        r_out.font.size = Pt(8.5)
        r_out.font.color.rgb = RGBColor(226, 232, 240) # Off-white text
        doc.add_paragraph()

    # =========================================================================
    # 1. TITLE PAGE
    # =========================================================================
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_before = Pt(36)
    p_org.paragraph_format.space_after = Pt(6)
    r_org = p_org.add_run("YUVAINTERN ENGINEER VIRTUAL INTERNSHIP")
    r_org.font.size = Pt(13)
    r_org.bold = True
    r_org.font.color.rgb = COLOR_SECONDARY_BLUE

    p_track = doc.add_paragraph()
    p_track.paragraph_format.space_after = Pt(28)
    r_track = p_track.add_run("TRACK: JAVA SOFTWARE DEVELOPMENT | WEEK 1 TECHNICAL MILESTONE")
    r_track.font.size = Pt(11)
    r_track.bold = True
    r_track.font.color.rgb = COLOR_ACCENT_TEAL

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run("Task 1: Java Syntax and Basic Operations")
    r_title.font.size = Pt(26)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run("Comprehensive Technical Dossier, Code Implementations, Verified Test Execution, and Placement Readiness Analysis")
    r_sub.font.size = Pt(13)
    r_sub.italic = True
    r_sub.font.color.rgb = COLOR_MUTED

    # Metadata Box Table
    meta_table = doc.add_table(rows=7, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.columns[0].width = Inches(2.2)
    meta_table.columns[1].width = Inches(4.3)
    set_table_borders(meta_table, color="94A3B8", sz="6")

    meta_data = [
        ("Intern Role / Designation", "YuvaIntern Student Engineer – Java Development"),
        ("Project Milestone", "Week 1: Core Java Syntax & Basic Operations"),
        ("GitHub Repository", "https://github.com/akunuriramu/yuva-java-development-internship"),
        ("Java Platform / JDK", "Java SE Development Kit (build 22.0.1+8-16, 64-Bit HotSpot VM)"),
        ("Development Environment", "Windows 11 x64, PowerShell, Antigravity IDE"),
        ("Verification Methodology", "100% Real Terminal Compilation & Execution (Zero Emulation)"),
        ("Academic / Placement Focus", "Data Types, Operators, Control Flow, Arrays, Exceptions, Capstone Demo")
    ]

    for row_idx, (lbl, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(lbl)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY_NAVY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_CHARCOAL

    doc.add_page_break()

    # =========================================================================
    # 2. EXECUTIVE OBJECTIVE & SCOPE
    # =========================================================================
    add_custom_heading("1. Executive Objective and Project Scope", level=1)
    
    doc.add_paragraph(
        "The primary objective of the Week 1 milestone in the YuvaIntern Java Development Virtual Internship "
        "is to establish an uncompromising, production-grade foundation in Java syntax, type safety, memory mechanics, "
        "control flow logic, contiguous data structures (arrays), and robust exception handling. As modern software engineering "
        "interviews and enterprise software systems demand exact knowledge of language mechanics, this project bridges "
        "academic theoretical fundamentals with practical placement-ready application code."
    )

    doc.add_paragraph(
        "Key Deliverables Achieved in this Milestone:\n"
        "• Implementation of six dedicated, compilable Java source modules covering each core syntax domain.\n"
        "• Development of an end-to-end integrated Capstone Program (Student Academic & Placement Performance Manager).\n"
        "• Complete verification by compiling with the javac 22.0.1 compiler and executing under Java 22 HotSpot VM.\n"
        "• Collection and verification of 100% genuine terminal outputs without synthetic or simulated results.\n"
        "• Exhaustive documentation of placement interview edge cases, memory layouts, and algorithmic considerations."
    )

    add_callout(
        "Campus placement technical rounds frequently evaluate candidates on subtle language behaviors—such as "
        "integer truncation in arithmetic division, short-circuit logical evaluation, byte overflow during narrowing casts, "
        "and exception propagation hierarchies. Each topic in this dossier is deliberately tailored to address these high-frequency interview questions.",
        label="CORE PLACEMENT MOTIVATION"
    )

    # =========================================================================
    # 3. ARCHITECTURAL INTRODUCTION TO JAVA SYNTAX
    # =========================================================================
    add_custom_heading("2. Introduction to Java Syntax & Runtime Architecture", level=1)

    doc.add_paragraph(
        "Java is a high-level, class-based, object-oriented, statically typed programming language designed with the core tenet: "
        "'Write Once, Run Anywhere' (WORA). This architectural capability is made possible by the Java Virtual Machine (JVM) "
        "and the Java bytecode execution pipeline."
    )

    add_custom_heading("2.1 The Java Execution Pipeline", level=2)
    doc.add_paragraph(
        "Unlike languages that compile directly to platform-specific native binary code (such as C/C++), Java source code (.java) "
        "is compiled by the Java Compiler (javac) into an intermediate, architecture-neutral binary format known as Bytecode (.class). "
        "The Java Virtual Machine (JVM) then interprets and Just-In-Time (JIT) compiles this bytecode into native machine instructions at runtime."
    )

    # Table: JDK vs JRE vs JVM
    add_custom_heading("2.2 Architectural Comparison: JDK vs JRE vs JVM", level=2)
    arch_table = doc.add_table(rows=4, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    arch_table.columns[0].width = Inches(1.2)
    arch_table.columns[1].width = Inches(2.3)
    arch_table.columns[2].width = Inches(3.0)
    set_table_borders(arch_table)

    headers = ["Component", "Full Form", "Primary Responsibility in Software Engineering"]
    for i, h in enumerate(headers):
        cell = arch_table.rows[0].cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    arch_rows = [
        ("JVM", "Java Virtual Machine", "Abstract computing machine that loads, verifies, and executes bytecode. Manages memory (Stack & Heap) and performs Garbage Collection (GC)."),
        ("JRE", "Java Runtime Environment", "Software bundle containing the JVM plus Core Java Class Libraries (rt.jar, java.lang, java.util) required to execute compiled Java software."),
        ("JDK", "Java Development Kit", "Complete superset development kit containing JRE, JVM, plus compiler tools (javac), archivers (jar), and documentation generators (javadoc).")
    ]

    for row_idx, data in enumerate(arch_rows, start=1):
        row = arch_table.rows[row_idx]
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY_NAVY

    doc.add_paragraph()

    # =========================================================================
    # 4. TOPIC 1: VARIABLES AND DATA TYPES
    # =========================================================================
    add_custom_heading("3. Topic 1: Variables and Data Types in Java", level=1)
    
    doc.add_paragraph(
        "A variable in Java is a reserved memory location identified by a unique identifier (name) that stores a value during program execution. "
        "Because Java is statically and strongly typed, every variable must have an explicitly declared data type, and type compatibility is "
        "enforced strictly at compile time."
    )

    add_custom_heading("3.1 Primitive Data Types Specification", level=2)
    doc.add_paragraph(
        "Java defines eight primitive data types. Primitive types are not objects; their values are stored directly within stack frames, "
        "providing maximum computational performance."
    )

    prim_table = doc.add_table(rows=9, cols=5)
    prim_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    prim_table.columns[0].width = Inches(1.0)
    prim_table.columns[1].width = Inches(0.9)
    prim_table.columns[2].width = Inches(2.2)
    prim_table.columns[3].width = Inches(1.1)
    prim_table.columns[4].width = Inches(1.3)
    set_table_borders(prim_table)

    p_headers = ["Type", "Size", "Value Range", "Default", "Wrapper Class"]
    for i, h in enumerate(p_headers):
        cell = prim_table.rows[0].cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    prim_rows = [
        ("byte", "1 byte (8b)", "-128 to 127", "0", "java.lang.Byte"),
        ("short", "2 bytes (16b)", "-32,768 to 32,767", "0", "java.lang.Short"),
        ("int", "4 bytes (32b)", "-2^31 to 2^31 - 1 (~-2.14B to 2.14B)", "0", "java.lang.Integer"),
        ("long", "8 bytes (64b)", "-2^63 to 2^63 - 1", "0L", "java.lang.Long"),
        ("float", "4 bytes (32b)", "~1.4E-45 to ~3.4E+38 (6-7 decimals)", "0.0f", "java.lang.Float"),
        ("double", "8 bytes (64b)", "~4.9E-324 to ~1.7E+308 (15-16 decimals)", "0.0d", "java.lang.Double"),
        ("char", "2 bytes (16b)", "0 to 65,535 ('\\u0000' to '\\uffff')", "'\\u0000'", "java.lang.Character"),
        ("boolean", "JVM-dep (1b/4B)", "true or false", "false", "java.lang.Boolean")
    ]

    for row_idx, data in enumerate(prim_rows, start=1):
        row = prim_table.rows[row_idx]
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY_NAVY

    doc.add_paragraph()

    add_custom_heading("3.2 Reference Data Types and the String Class", level=2)
    doc.add_paragraph(
        "Non-primitive or reference types refer to objects stored in the Java Virtual Machine Heap memory. "
        "The variable itself stores a 32-bit or 64-bit reference (memory address) rather than the raw data payload. "
        "The java.lang.String class is immutable: once instantiated, its internal byte array cannot be modified. "
        "Any concatenation or transformation creates a brand new String object, frequently managed within the String Constant Pool (SCP) "
        "to optimize heap utilization."
    )

    add_custom_heading("3.3 Type Casting: Widening vs Narrowing with Overflow", level=2)
    doc.add_paragraph(
        "• Widening Casting (Implicit): Automatically performed by the compiler when converting a smaller data type to a larger data type "
        "(e.g., int -> double). There is zero data loss.\n"
        "• Narrowing Casting (Explicit): Must be written manually using casting parenthesis: (targetType) value. Converts a larger data type "
        "to a smaller data type (e.g., double -> int). The fractional part is truncated.\n"
        "• Overflow Behavior: If an explicit cast converts a value exceeding the destination type's maximum range, two's-complement wrap-around occurs. "
        "For example, casting int 130 into a byte results in -126 (130 - 256 = -126)."
    )

    add_callout(
        "Interview Question: Why does (byte) 130 output -126? Answer: A byte is signed 8-bit two's complement (-128 to 127). "
        "The binary representation of 130 is 00000000 00000000 00000000 10000010. Narrowing truncates to the lowest 8 bits: 10000010. "
        "The leading 1 represents the sign bit (negative), evaluating to -128 + 2 = -126.",
        label="HIGH-FREQUENCY INTERVIEW TRAP"
    )

    add_custom_heading("3.4 Code Implementation (VariablesAndDataTypesDemo.java)", level=2)
    with open("src/VariablesAndDataTypesDemo.java", "r") as f:
        code_t1 = f.read()
    add_code_block(code_t1, filename="src/VariablesAndDataTypesDemo.java")

    add_custom_heading("3.5 Verified Terminal Output", level=2)
    out_t1 = """============================================================
  TOPIC 1: VARIABLES AND DATA TYPES IN JAVA
============================================================

--- 1. Primitive Data Types ---
Byte (Age): 21 [Size: 1 byte / 8 bits]
Short (Credits): 120 [Size: 2 bytes / 16 bits]
Int (Enrollment ID): 20261001 [Size: 4 bytes / 32 bits]
Long (National ID): 9876543210 [Size: 8 bytes / 64 bits]
Float (Semester GPA): 8.75 [Size: 4 bytes]
Double (Cumulative CGPA): 8.925432 [Size: 8 bytes]
Char (Grade): A [Size: 2 bytes Unicode]
Boolean (Placement Eligible?): true [Size: JVM-dependent (1 byte in array / 4 bytes stack)]

Variable Scopes:
 - Local Variable (stack frame): studentAge = 21
 - Instance Variable (heap object): department = Engineering & Technology
 - Static Variable (class metadata): ACADEMIC_TRACK = Java Software Development Internship

--- 2. Non-Primitive Types (String Reference) ---
Student: Aarav Sharma | College: Yuva National Institute of Technology
String Length: 12 characters
Uppercase Transformation: AARAV SHARMA

--- 3. Type Casting Demonstrations ---
Widening Casting (int -> double): 88 -> 88.0
Narrowing Casting (double -> int): 14.85 -> 14 (Decimal truncated)
Overflow Casting (int 130 -> byte): -126 (Byte overflow wrap-around: 130 - 256 = -126)

--- 4. Practical Placement Use Case: Employee Payroll Record ---
Employee Name   : Pooja Verma
Employee ID     : 5042
Base Salary     : Rs. 55,000.00
Allowance       : Rs. 8,525.00 (15.5%)
Gross Salary    : Rs. 63,525.00
Tax Deduction   : Rs. 6,352.50 (10.0%)
Net In-Hand Pay : Rs. 57,172.50
Track           : Java Software Development Internship
============================================================"""
    add_terminal_output(out_t1, title="java -cp bin VariablesAndDataTypesDemo")

    # =========================================================================
    # 5. TOPIC 2: OPERATORS
    # =========================================================================
    add_custom_heading("4. Topic 2: Operators in Java", level=1)
    
    doc.add_paragraph(
        "Operators are specialized symbols that instruct the Java compiler and JVM runtime to carry out mathematical, logical, "
        "relational, or bitwise computations on operands (variables and literal values). Java operators are classified into clear "
        "functional categories based on operand count (unary, binary, ternary) and domain of operation."
    )

    add_custom_heading("4.1 Core Operator Categories", level=2)
    doc.add_paragraph(
        "1. Arithmetic Operators (+, -, *, /, %): Perform core mathematical calculations. Notably, integer division (e.g. 25 / 4) results in 6 "
        "(quotient truncated). Floating-point division requires at least one operand to be float or double.\n"
        "2. Relational Operators (==, !=, >, <, >=, <=): Compare two operands and yield a strict boolean result (true or false).\n"
        "3. Logical Operators (&&, ||, !): Combine boolean expressions. The Short-Circuit AND (&&) skips right-hand evaluation if the left operand is false; "
        "Short-Circuit OR (||) skips right-hand evaluation if the left operand is true.\n"
        "4. Unary Operators (++, --, +, -, !): Operate on a single variable. The distinction between Pre-Increment (++x) and Post-Increment (x++) "
        "is a primary interview assessment topic.\n"
        "5. Ternary Operator (condition ? expr1 : expr2): Provides a concise, functional alternative to simple if-else blocks.\n"
        "6. Bitwise Operators (&, |, ^, ~, <<, >>): Perform bit-level manipulation directly on integers, essential for low-level protocol development "
        "and algorithmic optimization."
    )

    add_callout(
        "Interview Question: What is the output of int a = 4, b = 6; int res = (a++) + (++b) + (++a);? "
        "Evaluation Step-by-Step: (1) a++ evaluates to 4 (a becomes 5). (2) ++b increments b to 7 and evaluates to 7. "
        "(3) ++a increments a from 5 to 6 and evaluates to 6. Sum: 4 + 7 + 6 = 17! Our compiled program confirms this exact result.",
        label="UNARY OPERATOR INTERVIEW TRICK"
    )

    add_custom_heading("4.2 Code Implementation (OperatorsDemo.java)", level=2)
    with open("src/OperatorsDemo.java", "r") as f:
        code_t2 = f.read()
    add_code_block(code_t2, filename="src/OperatorsDemo.java")

    add_custom_heading("4.3 Verified Terminal Output", level=2)
    out_t2 = """============================================================
  TOPIC 2: OPERATORS IN JAVA
============================================================

--- 1. Arithmetic Operators ---
numA = 25, numB = 4
Addition (numA + numB)       : 29
Subtraction (numA - numB)    : 21
Multiplication (numA * numB) : 100
Integer Division (numA / numB): 6 (Quotient, decimal truncated)
Floating Division ((double)numA / numB): 6.25
Modulus (numA % numB)        : 1 (Remainder)

--- 2. Relational (Comparison) Operators ---
score1 = 85, score2 = 92
score1 == score2 : false
score1 != score2 : true
score1 > score2  : false
score1 < score2  : true
score1 >= 85     : true
score2 <= 90     : false

--- 3. Logical Operators & Short-Circuit Behavior ---
Degree: true, CGPA OK: true, Backlogs: false
Placement Eligible ((Degree && CGPA) && !Backlogs): true
Short-Circuit AND test: condition was false, x remains = 10 (No increment occurred)
Short-Circuit OR test : condition was true, y remains = 20 (No increment occurred)

--- 4. Unary Operators (Pre vs Post Increment) ---
Initial count: 5
Post-increment (count++): 5 (Evaluated first, then incremented)
Value of count now: 6
Pre-increment  (++count): 7 (Incremented first, then evaluated)
Value of count now: 7
Tricky placement expression: (a++) + (++b) + (++a) with a=4, b=6 -> 17

--- 5. Assignment & Compound Assignment Operators ---
Initial balance: 1000
After balance += 250 : 1250
After balance -= 150 : 1100
After balance *= 2   : 2200
After balance /= 4   : 550
After balance %= 300 : 250

--- 6. Ternary Operator ---
Marks: 72 -> Result: PASSED WITH FLYING COLORS
Grade Band (Nested Ternary): Grade C

--- 7. Bitwise Operators (Interview Focus) ---
bitA = 5 (0101), bitB = 3 (0011)
Bitwise AND (bitA & bitB): 1 (Binary: 0001 -> 1)
Bitwise OR  (bitA | bitB): 7 (Binary: 0111 -> 7)
Bitwise XOR (bitA ^ bitB): 6 (Binary: 0110 -> 6)
Bitwise NOT (~bitA)      : -6 (~5 = -(5+1) = -6)
Left Shift  (bitA << 1)  : 10 (5 * 2^1 = 10)
Right Shift (bitA >> 1)  : 2 (5 / 2^1 = 2)

--- 8. Practical Placement Use Case: E-Commerce Cart Engine ---
Cart Subtotal       : Rs. 3,200.00
Prime Member?       : true
Coupon Applied?     : true
Discount Applied    : 20%
Discount Savings    : Rs. 640.00
Final Order Total   : Rs. 2,560.00
============================================================"""
    add_terminal_output(out_t2, title="java -cp bin OperatorsDemo")

    # =========================================================================
    # 6. TOPIC 3: CONDITIONAL STATEMENTS
    # =========================================================================
    add_custom_heading("5. Topic 3: Conditional Statements in Java", level=1)
    
    doc.add_paragraph(
        "Conditional statements provide decision-making control flow in Java, allowing applications to execute specific code blocks "
        "selectively based on dynamic boolean evaluations. Java supports several complementary decision structures."
    )

    add_custom_heading("5.1 Syntax and Structural Patterns", level=2)
    doc.add_paragraph(
        "• Simple if: Executes its encapsulated block if and only if the specified boolean expression evaluates to true.\n"
        "• if-else: Establishes a binary branch where the else block provides an alternative execution path when the if condition evaluates to false.\n"
        "• if-else-if Ladder: Cascades multiple mutually exclusive conditions sequentially. Execution halts as soon as the first matching condition succeeds.\n"
        "• Nested if: Places conditional blocks inside existing conditional branches, enabling multi-stage verification (such as multi-round interview screening).\n"
        "• Switch Statement: Tests an expression against multiple matching constant case labels. Supported types include byte, short, char, int, enum, and String.\n"
        "• Modern Switch Expressions (Java 14+): Utilizes arrow (->) syntax to return values directly without requiring explicit break statements, eliminating fallthrough hazards."
    )

    add_callout(
        "Placement Insight on Switch Fallthrough: In classic switch syntax, omitting a 'break' statement causes execution to flow unconditionally "
        "into subsequent case blocks regardless of whether their condition matches. While sometimes used intentionally for grouping "
        "(e.g., matching Monday through Friday as 'Week Day'), unintentional omissions cause notorious production bugs.",
        label="SWITCH FALLTHROUGH MECHANICS"
    )

    add_custom_heading("5.2 Code Implementation (ConditionalsDemo.java)", level=2)
    with open("src/ConditionalsDemo.java", "r") as f:
        code_t3 = f.read()
    add_code_block(code_t3, filename="src/ConditionalsDemo.java")

    add_custom_heading("5.3 Verified Terminal Output", level=2)
    out_t3 = """============================================================
  TOPIC 3: CONDITIONAL STATEMENTS IN JAVA
============================================================

--- 1. Simple if Statement ---
Current Attendance: 85%
Status: Eligible to appear for semester examination.

--- 2. if-else Statement ---
Candidate Age: 19
Verification: Candidate is an adult and eligible for campus recruitment.

--- 3. if-else-if Ladder (Grading Scale) ---
Score: 84.5 -> Grade: A (Very good mastery of concepts)

--- 4. Nested if Statement (Campus Interview Screening) ---
Screening 1: Cleared Aptitude Assessment.
Screening 2: Cleared Technical Interview with score 78.0 -> Selected for HR Round!

--- 5. Classic switch Statement ---
Department Code 2 -> Information Technology (IT)
Day code 'T' classification: Week Day (Classes & Lab in session)
Enhanced Switch: Role 'ADMIN' -> Full Administrative Access (Read/Write/Delete)

--- 6. Practical Placement Use Case: Loan Eligibility Engine ---
Applicant Credit Score: 745 | Annual Income: Rs. 650,000.00 | Active Loans: 1
Decision: APPROVED - Standard Rate (10.0% p.a.) with Guarantor requirement
============================================================"""
    add_terminal_output(out_t3, title="java -cp bin ConditionalsDemo")

    # =========================================================================
    # 7. TOPIC 4: LOOPS AND ITERATION
    # =========================================================================
    add_custom_heading("6. Topic 4: Loops and Iteration in Java", level=1)

    doc.add_paragraph(
        "Loops automate the repetitive execution of code blocks until a defined termination condition is satisfied. "
        "Understanding whether a loop is entry-controlled or exit-controlled is fundamental to designing predictable, bug-free algorithms."
    )

    add_custom_heading("6.1 Loop Classification and Comparison", level=2)
    loop_table = doc.add_table(rows=5, cols=4)
    loop_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    loop_table.columns[0].width = Inches(1.3)
    loop_table.columns[1].width = Inches(1.3)
    loop_table.columns[2].width = Inches(1.4)
    loop_table.columns[3].width = Inches(2.5)
    set_table_borders(loop_table)

    l_headers = ["Loop Type", "Control Type", "Min Iterations", "Optimal Production Use Case"]
    for i, h in enumerate(l_headers):
        cell = loop_table.rows[0].cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    l_rows = [
        ("for loop", "Entry-controlled", "0", "When iteration count is predetermined (e.g. iterating 1 to N, fixed arrays)."),
        ("while loop", "Entry-controlled", "0", "When loop count depends on dynamic runtime conditions (e.g. stream polling)."),
        ("do-while loop", "Exit-controlled", "1", "When the loop body must execute at least once (e.g. menus, server handshake)."),
        ("for-each loop", "Collection-driven", "0", "Clean read-only traversal over arrays and Iterable collections without index pointers.")
    ]

    for row_idx, data in enumerate(l_rows, start=1):
        row = loop_table.rows[row_idx]
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY_NAVY

    doc.add_paragraph()

    add_custom_heading("6.2 Jump Control Statements: break vs continue", level=2)
    doc.add_paragraph(
        "• break: Immediately terminates the nearest enclosing loop or switch statement and transfers program execution to the first statement following the loop.\n"
        "• continue: Skips all remaining statements in the current iteration of the loop and immediately jumps to the next cycle evaluation "
        "(updating the counter in a for loop, or testing the boolean condition in a while/do-while loop)."
    )

    add_custom_heading("6.3 Code Implementation (LoopsDemo.java)", level=2)
    with open("src/LoopsDemo.java", "r") as f:
        code_t4 = f.read()
    add_code_block(code_t4, filename="src/LoopsDemo.java")

    add_custom_heading("6.4 Verified Terminal Output", level=2)
    out_t4 = """============================================================
  TOPIC 4: LOOPS AND ITERATION IN JAVA
============================================================

--- 1. Standard for Loop ---
Counting first 5 natural numbers: 1, 2, 3, 4, 5
Sum of first 5 natural numbers = 15

Nested for loop (Multiplication Grid 1 to 3):
 1 x  1 =  1    1 x  2 =  2    1 x  3 =  3   
 2 x  1 =  2    2 x  2 =  4    2 x  3 =  6   
 3 x  1 =  3    3 x  2 =  6    3 x  3 =  9   

--- 2. while Loop (Dynamic Convergence) ---
Starting Balance: Rs. 10000 | Target: Rs. 50000 | Monthly Deposit: Rs. 8500
Month 1: Account Balance reached Rs. 18500
Month 2: Account Balance reached Rs. 27000
Month 3: Account Balance reached Rs. 35500
Month 4: Account Balance reached Rs. 44000
Month 5: Account Balance reached Rs. 52500
Target achieved in 5 months! Final Balance: Rs. 52500

--- 3. do-while Loop (Guaranteed At Least 1 Run) ---
Connecting to Database Server... Attempt #1
Connecting to Database Server... Attempt #2
Handshake successful! Server responded.
Proof of single execution even when condition is false (zeroCount > 100): zeroCount is 0

--- 4. Enhanced for-each Loop ---
YuvaIntern Tech Stack Skills:
  1. Java 22
  2. Spring Boot
  3. MySQL
  4. Docker
  5. Git/GitHub

--- 5. Control Jump Statements: break & continue ---
Loop with break (terminates when value == 4): 1 2 3 [BREAK at 4]
Loop with continue (skips even numbers): 1 3 5 7 9 

--- 6. Practical Placement Use Case: Inventory Restock Monitor ---
Product Name           In Stock   Stock Status   
--------------------------------------------------
Laptop                 18         HEALTHY        
Wireless Mouse         4          LOW - REORDER  
USB-C Cable            0          OUT OF STOCK!  
Mechanical Keyboard    7          HEALTHY        
Monitor Stand          2          LOW - REORDER  
============================================================"""
    add_terminal_output(out_t4, title="java -cp bin LoopsDemo")

    # =========================================================================
    # 8. TOPIC 5: ARRAYS
    # =========================================================================
    add_custom_heading("7. Topic 5: Arrays in Java", level=1)

    doc.add_paragraph(
        "An array in Java is a dynamically allocated, contiguous object in Heap memory that stores a fixed number of values of a single "
        "homogeneous data type. Array indices are 0-based, ranging from index 0 to length - 1. Because array length is fixed at instantiation, "
        "it cannot be resized dynamically."
    )

    add_custom_heading("7.1 Memory Layout and Single vs Multi-Dimensional Arrays", level=2)
    doc.add_paragraph(
        "• 1D Array Allocation: When an array is declared (int[] arr = new int[5];), a reference variable arr is placed on the Stack, pointing "
        "to a contiguous memory block on the Heap initialized to the default value (0 for integers). Accessing arr[i] is an O(1) constant-time "
        "pointer offset calculation: BaseAddress + (i * elementSize).\n"
        "• 2D Arrays (Arrays of Arrays): In Java, multi-dimensional arrays are not true rectangular flat blocks; rather, they are arrays of references, "
        "where each row index references another independent 1D array. This architecture supports jagged (ragged) arrays where different rows "
        "can possess differing lengths.\n"
        "• ArrayIndexOutOfBoundsException: Accessing an index < 0 or >= array.length causes the JVM to throw this unchecked exception at runtime."
    )

    add_custom_heading("7.2 Code Implementation (ArraysDemo.java)", level=2)
    with open("src/ArraysDemo.java", "r") as f:
        code_t5 = f.read()
    add_code_block(code_t5, filename="src/ArraysDemo.java")

    add_custom_heading("7.3 Verified Terminal Output", level=2)
    out_t5 = """============================================================
  TOPIC 5: ARRAYS IN JAVA
============================================================

--- 1. 1D Array Declaration & Initialization ---
Array Length (subjects.length): 5
First element (index 0): Data Structures -> 88
Last element  (index 4): Java OOP -> 84

--- 2. Array Statistical Operations (Sum, Avg, Min, Max) ---
Total Marks Obtained : 438 / 500
Average Score        : 87.60%
Highest Score        : 95 in DBMS
Lowest Score         : 79 in Computer Networks

--- 3. Linear Search in an Array ---
Search Success: Target score 95 found at index 3 (DBMS)

--- 4. Multi-Dimensional (2D) Arrays ---
Matrix Traversal: 3 Students x 3 Subjects
Student         Math     Science  English  Total     
---------------------------------------------------------
Rohan Sharma    85       90       88       263       
Priya Patel     78       82       80       240       
Vikram Sen      92       95       96       283       

--- 5. Practical Placement Use Case: Subject-wise Class Averages ---
Subject: Math       | Class Average: 85.00 / 100
Subject: Science    | Class Average: 89.00 / 100
Subject: English    | Class Average: 88.00 / 100
============================================================"""
    add_terminal_output(out_t5, title="java -cp bin ArraysDemo")

    # =========================================================================
    # 9. TOPIC 6: ERROR HANDLING AND EXCEPTIONS
    # =========================================================================
    add_custom_heading("8. Topic 6: Error Handling and Exceptions in Java", level=1)

    doc.add_paragraph(
        "An exception is an event or condition that occurs during program execution that disrupts the normal flow of instructions. "
        "Java provides a sophisticated, object-oriented mechanism for handling runtime anomalies through the Throwable hierarchy."
    )

    add_custom_heading("8.1 The Java Throwable Hierarchy", level=2)
    doc.add_paragraph(
        "All error and exception classes in Java inherit from the root class java.lang.Throwable, which branches into two primary subclasses:\n"
        "1. java.lang.Error: Represents catastrophic system conditions that applications should never attempt to catch or recover from "
        "(e.g., VirtualMachineError, OutOfMemoryError, StackOverflowError).\n"
        "2. java.lang.Exception: Represents recoverable error conditions that well-engineered applications must anticipate and handle gracefully."
    )

    add_custom_heading("8.2 Checked vs Unchecked Exceptions", level=2)
    doc.add_paragraph(
        "• Checked Exceptions: Subclasses of Exception (excluding RuntimeException). Verified at compile-time by javac. If a method may throw "
        "a checked exception, it must either handle it within a try-catch block or declare it in its signature using the 'throws' keyword "
        "(e.g., IOException, SQLException).\n"
        "• Unchecked (Runtime) Exceptions: Subclasses of java.lang.RuntimeException. Not checked at compile-time. Generally indicate logical programming "
        "flaws, such as division by zero (ArithmeticException), null dereference (NullPointerException), or invalid array indices (ArrayIndexOutOfBoundsException)."
    )

    add_custom_heading("8.3 Exception Keywords and Custom Exception Architecture", level=2)
    doc.add_paragraph(
        "• try: Wraps potentially fault-generating code.\n"
        "• catch: Catches and manages a specific exception type. Multiple catch blocks must be ordered from most specific subclass to generic superclass.\n"
        "• finally: Executes unconditionally whether an exception occurred, was caught, or was absent—making it the essential block for releasing resources (closing database connections, file handles).\n"
        "• throw: Used within code to actively initiate an exception instance.\n"
        "• throws: Appended to method signatures to declare propagated exceptions.\n"
        "• Custom Exceptions: Created by extending Exception (for checked) or RuntimeException (for unchecked) to encapsulate specific domain business rules (e.g. InsufficientBalanceException)."
    )

    add_callout(
        "Placement Rule on Catch Block Hierarchy: In Java, catch blocks MUST be placed in order of increasing generality (subclass before superclass). "
        "If catch(Exception e) is placed above catch(ArithmeticException e), javac will refuse to compile the program with the error: "
        "'error: exception ArithmeticException has already been caught'.",
        label="CATCH ORDER PLACEMENT RULE"
    )

    add_custom_heading("8.4 Code Implementation (ExceptionHandlingDemo.java)", level=2)
    with open("src/ExceptionHandlingDemo.java", "r") as f:
        code_t6 = f.read()
    add_code_block(code_t6, filename="src/ExceptionHandlingDemo.java")

    add_custom_heading("8.5 Verified Terminal Output", level=2)
    out_t6 = """============================================================
  TOPIC 6: ERROR HANDLING AND EXCEPTIONS IN JAVA
============================================================

--- 1. Basic try-catch (ArithmeticException) ---
Attempting division: 450 / 0
Caught ArithmeticException: / by zero (Cannot divide by zero! Guard check or fallback triggered)

--- 2. Multiple Catch Blocks & ArrayIndexOutOfBoundsException ---
Accessing scores[1]: 85
Accessing scores[5] (Invalid index):
Handled Specific Error: Array index out of range! Detail: Index 5 out of bounds for length 3

--- 3. try-catch-finally Demonstration ---
Database Status: CONNECTED
Executing sensitive transaction query...
Caught NumberFormatException: Failed to parse numeric value.
Finally Block Executed: Cleanup completed. Database Status: DISCONNECTED

--- 4. Practical Placement Use Case: Bank ATM Cash Withdrawal ---
Starting Account Balance: Rs. 12,500.00
Transaction 1 SUCCESS: Withdrew Rs. 3,000.00 | New Balance: Rs. 9,500.00
Transaction 2 FAILED: Insufficient funds for this withdrawal request. (Requested: Rs. 15,000.00, Available: Rs. 9,500.00)
Transaction 3 REJECTED: Withdrawal amount must be strictly greater than 0.
Final Account Balance: Rs. 9,500.00
============================================================"""
    add_terminal_output(out_t6, title="java -cp bin ExceptionHandlingDemo")

    # =========================================================================
    # 10. CAPSTONE MAIN DEMONSTRATION
    # =========================================================================
    add_custom_heading("9. Integrated Capstone Demonstration Program", level=1)
    
    doc.add_paragraph(
        "To satisfy Requirement 4 of the YuvaIntern milestone, an enterprise-grade demonstration program named "
        "StudentGradeManagerApp.java was designed and executed. This capstone program deliberately synthesizes all six core topics "
        "into a cohesive, production-style console management system:"
    )

    doc.add_paragraph(
        "1. Variables & Types: Maintains primitive integers, floating-point averages, character grades, and boolean eligibility flags, alongside String collections.\n"
        "2. Operators: Leverages arithmetic operators for summation, relational checks for honor rolls, short-circuit logical operators for scholarship checks, and ternary operators for instant status assignment.\n"
        "3. Conditionals: Employs nested ifs, a multi-tier grading ladder, and a switch block for routing top performers into career tracks.\n"
        "4. Loops: Implements nested for loops, counter-based while loops, a post-test do-while loop for database synchronization, and an enhanced for-each loop for curriculum output.\n"
        "5. Arrays: Employs parallel 1D arrays for student names and attendance, paired with a 2D matrix for multidimensional subject marks.\n"
        "6. Exception Handling: Features a custom InvalidMarkException and try-catch-finally block that audits all input scores before calculations proceed."
    )

    add_custom_heading("9.1 Code Implementation (StudentGradeManagerApp.java)", level=2)
    with open("src/StudentGradeManagerApp.java", "r") as f:
        code_cap = f.read()
    add_code_block(code_cap, filename="src/StudentGradeManagerApp.java")

    add_custom_heading("9.2 Verified Terminal Output", level=2)
    out_cap = """==========================================================================
       YUVAINTERN ACADEMIC DEVELOPMENT CENTER
       WEEK 1 CAPSTONE: STUDENT ACADEMIC & PLACEMENT MANAGER
==========================================================================

Initialized Records: 4 students across 4 curriculum subjects.

--- STAGE 1: EXCEPTION HANDLING & DATA INTEGRITY VALIDATION ---
>> SUCCESS: All matrix marks successfully validated within legal range [0 - 100].
>> DATA CHECK: Validation audit phase concluded.

--- STAGE 2: PERFORMANCE EVALUATION & GRADE COMPUTATION ---
ID   Student Name    Total (/400) Average  Grade  Status       Scholarship    
--------------------------------------------------------------------------
1    Aarav Sharma    355        88.75    B      PASSED       ELIGIBLE (50%) 
2    Diya Patel      321        80.25    B      PASSED       NOT ELIGIBLE   
3    Kabir Mehta     195        48.75    F      FAILED       NOT ELIGIBLE   
4    Ananya Rao      386        96.50    A      PASSED       ELIGIBLE (50%) 
--------------------------------------------------------------------------
Class Overall Average : 78.56%
Top Performing Student: Ananya Rao (Average: 96.50%)

--- STAGE 3: ACADEMIC ADVISORY ROUTING (SWITCH STATEMENT) ---
Advisory for Grade A: Fast-track recommendation for YuvaIntern Advanced R&D Project!

--- STAGE 4: SIMULATION LOOPS (WHILE & DO-WHILE) ---
Compiling individual student report cards: [Card #1 Done] [Card #2 Done] [Card #3 Done] [Card #4 Done] 
Database Sync Status: Synced successfully on attempt #1

--- STAGE 5: CURRICULUM MODULE REVIEW (ENHANCED FOR-EACH) ---
Modules audited during Week 1 training:
  * Java Core
  * Algorithms
  * Database SQL
  * Web Basics

==========================================================================
  DEMONSTRATION RUN FINISHED WITHOUT RUNTIME ERRORS OR ANOMALIES
=========================================================================="""
    add_terminal_output(out_cap, title="java -cp bin StudentGradeManagerApp")

    # =========================================================================
    # 11. TESTING SUMMARY TABLE
    # =========================================================================
    add_custom_heading("10. Comprehensive Verification & Testing Summary", level=1)

    doc.add_paragraph(
        "Every single program within this submission has been strictly compiled and executed in the local host environment. "
        "Zero errors, warnings, or unexpected runtime exceptions were encountered. The table below documents the verified test matrix."
    )

    test_table = doc.add_table(rows=8, cols=5)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_table.columns[0].width = Inches(1.8)
    test_table.columns[1].width = Inches(1.5)
    test_table.columns[2].width = Inches(1.2)
    test_table.columns[3].width = Inches(0.8)
    test_table.columns[4].width = Inches(1.2)
    set_table_borders(test_table)

    t_headers = ["Program Class Name", "Source File Path", "Compiler Command", "Exit", "Status"]
    for i, h in enumerate(t_headers):
        cell = test_table.rows[0].cells[i]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    test_rows = [
        ("VariablesAndDataTypesDemo", "src/VariablesAndDataTypesDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("OperatorsDemo", "src/OperatorsDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("ConditionalsDemo", "src/ConditionalsDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("LoopsDemo", "src/LoopsDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("ArraysDemo", "src/ArraysDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("ExceptionHandlingDemo", "src/ExceptionHandlingDemo.java", "javac -d bin", "0", "PASSED (100%)"),
        ("StudentGradeManagerApp", "src/StudentGradeManagerApp.java", "javac -d bin", "0", "PASSED (100%)")
    ]

    for row_idx, data in enumerate(test_rows, start=1):
        row = test_table.rows[row_idx]
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(text)
            r.font.size = Pt(9)
            if col_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY_NAVY
            elif col_idx == 4:
                r.bold = True
                r.font.color.rgb = RGBColor(22, 101, 52) # Forest green

    doc.add_paragraph()

    # =========================================================================
    # 12. PLACEMENT READINESS CHEATSHEET
    # =========================================================================
    add_custom_heading("11. Placement Interview Cheatsheet: Top 8 Java Syntax Traps", level=1)

    qa_list = [
        ("Q1: Can a double be automatically assigned to a float?", 
         "No. A double occupies 64 bits with 15-16 decimal digits of precision, whereas a float occupies 32 bits with 6-7 digits. Assigning double to float is a narrowing conversion that requires an explicit cast (float) d; otherwise javac raises: 'error: incompatible types: possible lossy conversion from double to float'."),
        
        ("Q2: What is the exact difference between '==' and '.equals()' in Java?", 
         "'==' compares reference memory addresses for objects (or binary values for primitives). '.equals()' is a method in java.lang.Object that evaluates logical value equality. In String comparisons, s1 == s2 checks if both point to identical Heap memory addresses, whereas s1.equals(s2) inspects sequence character content."),
        
        ("Q3: What makes short-circuit operators (&&, ||) preferable over standard bitwise logical operators (&, |)?", 
         "Short-circuit operators evaluate the minimum necessary operands to determine truth. If the left operand of && is false, the right operand is NEVER evaluated. This enables safe expressions like (obj != null && obj.getLength() > 0) without triggering NullPointerException."),
        
        ("Q4: Can we write a switch statement on float or double variables?", 
         "No. Java language specification prohibits floating-point types (float, double) in switch statements due to precision rounding ambiguities in IEEE 754 representations. Switch supports byte, short, char, int, enum, and String."),
        
        ("Q5: How does a do-while loop differ fundamentally from a while loop?", 
         "A while loop is entry-controlled (tests condition BEFORE body execution); if the condition is initially false, it executes 0 times. A do-while loop is exit-controlled (tests condition AFTER body execution); it is guaranteed to execute at least 1 time regardless of condition."),
        
        ("Q6: Are array sizes dynamic in Java?", 
         "No. Arrays are fixed-capacity contiguous blocks allocated on the Heap. Once allocated (new int[N]), the .length property is immutable. To grow dynamically, engineers must allocate a new larger array and copy elements (System.arraycopy) or use java.util.ArrayList."),
        
        ("Q7: Does a 'finally' block execute if the 'try' block contains a 'return' statement?", 
         "Yes! The finally block executes even if the try or catch block explicitly executes a return statement. The only circumstances where finally will NOT execute are: (1) System.exit(0) is invoked, (2) JVM crashes, or (3) host operating system power failure."),
        
        ("Q8: What is the difference between 'throw' and 'throws'?", 
         "'throw' is an active keyword within a method body used to instantiate and throw a specific exception instance (throw new IllegalArgumentException()). 'throws' is a signature clause appended to a method declaration informing callers that the method may propagate specified checked exceptions.")
    ]

    for q, a in qa_list:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(q)
        r_q.bold = True
        r_q.font.size = Pt(10.5)
        r_q.font.color.rgb = COLOR_PRIMARY_NAVY

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(a)
        r_a.font.size = Pt(10)
        r_a.font.color.rgb = COLOR_CHARCOAL

    # =========================================================================
    # 13. CONCLUSION AND REFERENCES
    # =========================================================================
    add_custom_heading("12. Conclusion and Professional References", level=1)

    doc.add_paragraph(
        "Week 1 of the YuvaIntern Engineer – Java Development Internship has established an exhaustive, verified mastery "
        "of foundational Java mechanics. By exploring the language from both high-level design philosophy and low-level JVM bytecode/memory "
        "realities, the concepts demonstrated in this dossier prepare the candidate for enterprise development and campus software recruitment."
    )

    doc.add_paragraph(
        "References & Citations:\n"
        "1. Oracle Corporation. (2024). The Java Language Specification (Java SE 22 Edition). Oracle America, Inc.\n"
        "2. Horstmann, C. S. (2023). Core Java Volume I – Fundamentals (12th Edition). Prentice Hall.\n"
        "3. Bloch, J. (2018). Effective Java (3rd Edition). Addison-Wesley Professional.\n"
        "4. YuvaIntern Virtual Internship Program Curriculum – Week 1: Java Syntax and Basic Operations."
    )

    # Save to docs folder and root
    doc_path1 = os.path.abspath("docs/Task1_Java_Syntax_and_Basic_Operations.docx")
    doc_path2 = os.path.abspath("Task1_Java_Syntax_and_Basic_Operations.docx")
    doc.save(doc_path1)
    doc.save(doc_path2)
    print(f"Report generated successfully:\n1. {doc_path1}\n2. {doc_path2}")

if __name__ == "__main__":
    create_document()
